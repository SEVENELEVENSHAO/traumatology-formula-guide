from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pptx import Presentation


DEFAULT_SOURCES = [
    Path(r"E:\Downloads\ASUS\Yr 4 external medicine 414C student copy.pptx"),
    Path(r"E:\Downloads\ASUS\formula runs 2026.docx"),
    Path(r"E:\Downloads\ASUS\Formulas for Final Exam.docx"),
]

FORMULA_MARKER = re.compile(
    r"(?i)\b(?:tang|san|wan|yin|dan|jian|gao|pill|powder|formula|decoction)\b"
)
CONTEXT_MARKER = re.compile(
    r"(?i)\b(?:syndrome|stage|treatment|fracture|stasis|deficiency|damp|cold|heat|pain|trauma|swelling|tendon|bone)\b"
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def compact(value: str) -> str:
    return " ".join(value.replace("\u000b", " ").split())


def source_meta(path: Path, source_id: str, kind: str) -> dict:
    stat = path.stat()
    return {
        "id": source_id,
        "title": path.stem,
        "kind": kind,
        "language": "English with pinyin",
        "path": str(path),
        "sha256": sha256(path),
        "sizeBytes": stat.st_size,
        "modifiedUtc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
        "rightsScope": "public redistribution confirmed by user",
        "reviewStatus": "pending",
    }


def extract_docx(path: Path, source_id: str) -> tuple[dict, list[dict]]:
    document = Document(path)
    records: list[dict] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = compact(paragraph.text)
        if not text:
            continue
        records.append(
            {
                "id": f"{source_id}:p{index}",
                "sourceId": source_id,
                "locator": {"kind": "paragraph", "index": index},
                "style": paragraph.style.name,
                "text": text,
                "formulaRelevant": bool(FORMULA_MARKER.search(text)),
                "extractionMethod": "python-docx text layer",
            }
        )
    return source_meta(path, source_id, "docx"), records


def slide_text(slide) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        value = getattr(shape, "text", "")
        value = compact(value)
        if value:
            parts.append(value)
    return " / ".join(parts)


def extract_pptx(path: Path, source_id: str) -> tuple[dict, list[dict], dict]:
    presentation = Presentation(path)
    all_slides = [slide_text(slide) for slide in presentation.slides]
    relevant: set[int] = set()
    for index, text in enumerate(all_slides):
        if FORMULA_MARKER.search(text):
            relevant.add(index)
            if index > 0 and CONTEXT_MARKER.search(all_slides[index - 1]):
                relevant.add(index - 1)

    records: list[dict] = []
    for index in sorted(relevant):
        slide = presentation.slides[index]
        text = all_slides[index]
        picture_count = sum(1 for shape in slide.shapes if int(shape.shape_type) == 13)
        records.append(
            {
                "id": f"{source_id}:s{index + 1}",
                "sourceId": source_id,
                "locator": {"kind": "slide", "index": index + 1},
                "text": text,
                "pictureCount": picture_count,
                "formulaRelevant": bool(FORMULA_MARKER.search(text)),
                "needsVisualReview": picture_count > 0 and len(text) < 80,
                "extractionMethod": "python-pptx text layer",
            }
        )
    report = {
        "totalSlides": len(presentation.slides),
        "includedSlides": len(records),
        "excludedSlides": len(presentation.slides) - len(records),
        "visualReviewSlides": [
            item["locator"]["index"] for item in records if item["needsVisualReview"]
        ],
    }
    return source_meta(path, source_id, "pptx"), records, report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("sources", nargs="*", type=Path, default=DEFAULT_SOURCES)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "src" / "data" / "raw-sources.json",
    )
    args = parser.parse_args()

    source_ids = ["external-medicine-414c", "formula-runs-2026", "final-exam-formulas"]
    sources: list[dict] = []
    records: list[dict] = []
    reports: dict[str, dict] = {}
    for source_id, path in zip(source_ids, args.sources, strict=True):
        if not path.exists():
            raise FileNotFoundError(path)
        if path.suffix.lower() == ".docx":
            meta, extracted = extract_docx(path, source_id)
        elif path.suffix.lower() == ".pptx":
            meta, extracted, report = extract_pptx(path, source_id)
            reports[source_id] = report
        else:
            raise ValueError(f"Unsupported source: {path}")
        sources.append(meta)
        records.extend(extracted)

    payload = {
        "schemaVersion": 1,
        "generatedUtc": datetime.now(timezone.utc).isoformat(),
        "sources": sources,
        "records": records,
        "extractionReport": reports,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"sources": len(sources), "records": len(records), "output": str(args.output)}))


if __name__ == "__main__":
    main()
